#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сбор данных о людях "моей округи" из исповедных ведомостей и ревизских сказок Яндекс Архива (версия 1.1).

РЕКОНСТРУКЦИЯ. Оригинальный файл версии 1.1 не сохранился. Скрипт восстановлен обратно
из расширения yandex-archive-extension 1.1.0 (logic.js + page-scripts.js + app.js): логика разбора
перенесена из JavaScript в Python один в один. Оболочка (вход в Яндекс, запуск браузера, окна выбора
файлов, запись Excel) в расширении другая, поэтому взята из версии 1.3.

Что нового по сравнению с 1.0:
  * два типа обработки: «Исповедные ведомости» и «Ревизские сказки» (параметры — в словаре MODES);
    ревизские: доп. заголовки (починок, выселок, слобода, хутор, пустошь, погост), номер двора
    («Двор № 5», «№ 12»), два возраста в строке (колонки «Возраст: 1-е число» / «Возраст: 2-е число»),
    допускаются короткие записи;
  * текст берётся только из окна распознанного текста справа от окна со сканом (автоопределение);
    дата документа — строка выше панели инструментов окна со сканом;
    вместо ручной калибровки мышью в расширении здесь — настройки TEXT_SELECTOR / DATE_SELECTOR
    (CSS-селекторы; пусто = автоопределение);
  * служебные строки (название документа, уезд, приход...) отбрасываются;
  * если окно текста не найдено на 3 страницах подряд, работа останавливается (MAX_MISSES);
  * новые колонки Excel: «Тип документа», «Дата документа (со страницы)», «Двор №».

Диапазон страниц берётся РОВНО из ссылки. Остановка: Ctrl+C — собранное сохраняется.

Установка:
  python3 -m pip install playwright beautifulsoup4 openpyxl
  (браузер: запущенный Chrome с портом 9222 или установленный Google Chrome)
"""

import os
import random
import re
import sys
import time
import tkinter as tk
from tkinter import filedialog, messagebox
from urllib.parse import parse_qs, urlparse

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright

# ----------------------------------------------------------------------------
# НАСТРОЙКИ
# ----------------------------------------------------------------------------
ARCHIVE_HOSTS = ("ya.ru", "yandex.ru")      # равнозначные домены
WORK_HOST = "https://ya.ru"
PAGE_URL_TEMPLATE = WORK_HOST + "/archive/catalog/{file_id}/{page}"

HEADER_MAX_LEN = 70    # строка короче этого может быть заголовком населённого пункта
BIG_RANGE = 300        # предупреждение о слишком большом диапазоне
MAX_MISSES = 3         # подряд страниц без окна текста — остановка

DELAY_RANGE = (1.0, 2.5)
PAGE_TIMEOUT_MS = 45_000
MAX_RETRIES = 2
RENDER_WAIT_MS = 15_000
DEBUG_HTML = True                      # сохранить HTML первой страницы (first_page.html)
DEBUG_DIR = "debug_yandex_archive"

# Ручная настройка окон (аналог кнопок «Указать окно текста / строку даты вручную» в расширении).
# CSS-селектор; пустая строка — автоопределение.
TEXT_SELECTOR = ""
DATE_SELECTOR = ""

# ------------------------------------------------------------------
# РЕЖИМЫ ОБРАБОТКИ. Здесь меняются параметры сбора под тип документа.
# ------------------------------------------------------------------
MODES = {
    "confession": {
        "id": "confession",
        "label": "Исповедные ведомости",
        "extra_prefixes": [],       # слова-префиксы заголовка НП (кроме "сел…" и "дер…")
        "min_line_len": 6,          # минимальная длина строки с человеком
        "two_ages": False,          # в строке один возраст
        "household": False,         # номера дворов не отслеживаем
        "default_span": 15,         # если в ссылке нет sheet_page_to: from .. from+default_span
        "description": ("Заголовки населённых пунктов: село / деревня. В строке один возраст. "
                        "Люди собираются по строкам под заголовком вашего населённого пункта."),
    },
    "revision": {
        "id": "revision",
        "label": "Ревизские сказки",
        "extra_prefixes": ["починок", "починк", "выселок", "выселк", "слобод", "хутор", "пустош", "погост"],
        "min_line_len": 5,          # записи короче: «сын Иван 5»
        "two_ages": True,           # обычно два возраста: по прошлой ревизии и ныне
        "household": True,          # отслеживаем «Двор № 5» / «№ 12»
        "default_span": 15,
        "description": ("Заголовки: также починок, выселок, слобода, хутор и т. п. Отслеживается номер двора "
                        "(«Двор № 5», «№ 12»). В строке обычно два возраста — они выводятся в отдельные колонки. "
                        "Допускаются более короткие строки."),
    },
}
DEFAULT_MODE = "confession"
# ----------------------------------------------------------------------------


def get_mode(mode_id):
    return MODES.get(mode_id) or MODES[DEFAULT_MODE]


# строки-заголовки документа, которые не являются людьми
TITLE_NOISE = re.compile(r"(?:ревизск|исповедн|ведомост|сказк|уезд|губерн|волост|приход)", re.I)

# Границы слов (в Python \w понимает кириллицу)
WB = r"(?<!\w)"
WE = r"(?!\w)"


def ci(w: str) -> str:
    """Слово -> регулярка, нечувствительная к регистру (без флага re.I, чтобы [А-Я] оставалось строгим)."""
    return "".join(ch if ch.upper() == ch else f"[{ch.upper()}{ch}]" for ch in w)


# Населённые пункты округи: название -> основа (по нормализованному тексту)
SETTLEMENT_TAILS = {
    "село Троицкое":     r"троицк",
    "деревня Острогина": r"острогин",
    "деревня Черепкова": r"черепков",
    "деревня Рублева":   r"рублев",
    "деревня Луки":      r"лук(?:и)?" + WE,
    "деревня Мякинина":  r"м[яе]кинин",
}

_compiled = {}


def compile_mode(mode_id):
    m = get_mode(mode_id)
    if m["id"] in _compiled:
        return _compiled[m["id"]]
    words = "|".join(["сел", "дер"] + m["extra_prefixes"])
    words_ci = "|".join(ci(w) for w in ["сел", "дер"] + m["extra_prefixes"])

    # Обязателен префикс (в т.ч. "дер." / "д." / "с."). Применяется к нормализованному тексту.
    prefix = rf"(?:{WB}(?:{words})\w*\.?|{WB}[дс]\.)\s*[\[\(«\"]?\s*"
    settlements = {name: re.compile(prefix + tail) for name, tail in SETTLEMENT_TAILS.items()}

    # "Деревня"/"Дер." с любой буквы; сокращения "д."/"с." только строчные,
    # чтобы не принять за заголовок инициалы ("Д. Иванов").
    other_header = re.compile(
        rf"(?:{WB}(?:{words_ci})\w*\.?|{WB}[дс]\.)\s*[\[\(«\"]?\s*([А-ЯЁІѢ][а-яёіѣъь]{{3,}})")
    # строка, состоящая только из слова-префикса ("Деревня" / "Село" / "д.")
    prefix_only = re.compile(rf"^\s*(?:(?:{words_ci})\w*|[дсДС])\.?\s*$")

    # маркеры двора: "Двор № 5 ..." или строка "№ 12"
    dvor = ci("двор")
    hh_a = re.compile(rf"^\s*(?:{dvor}\w*|[Дд][Вв])\.?\s*(?:№|N|n)?\s*(\d{{1,3}})(?!\w)")
    hh_b = re.compile(r"^\s*(?:№|N)\s*(\d{1,3})\s*[.:]?\s*$")

    _compiled[m["id"]] = dict(settlements=settlements, other_header=other_header,
                              prefix_only=prefix_only, hh_a=hh_a, hh_b=hh_b)
    return _compiled[m["id"]]


def norm(s: str) -> str:
    s = s.lower()
    for a, b in (("ё", "е"), ("ѣ", "е"), ("і", "и"), ("ѵ", "и"), ("ѳ", "ф")):
        s = s.replace(a, b)
    return s.replace("ъ", "").replace("ь", "")


# ============================================================================
# JS, выполняемый внутри страницы архива (page-scripts.js): окно распознанного текста
# справа от окна со сканом + строка даты документа выше окна со сканом
# ============================================================================
EXTRACT_JS = r"""
async (opts) => {
  const textSel = opts.textSel || '';
  const dateSel = opts.dateSel || '';
  const maxWait = opts.maxWait || 15000;

  const rectOf = (e) => e.getBoundingClientRect();
  const visible = (e) => {
    const cs = getComputedStyle(e);
    if (cs.display === 'none' || cs.visibility === 'hidden') return false;
    const r = rectOf(e);
    return r.width > 0 && r.height > 0;
  };
  const clean = (s) => (s || '').replace(/[ \t ]+/g, ' ').replace(/\s*\n\s*/g, '\n').trim();
  const MONTHS = /январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|октябр|ноябр|декабр/i;

  // окно со сканом: самый большой canvas/img/svg
  const findViewer = () => {
    const list = [...document.querySelectorAll('canvas, img, svg')]
      .filter((e) => visible(e))
      .map((e) => ({ e, r: rectOf(e) }))
      .filter((x) => x.r.width > 200 && x.r.height > 200)
      .sort((a, b) => b.r.width * b.r.height - a.r.width * a.r.height);
    if (list.length) return list[0].r;
    return null;
  };

  const bySelector = (sel) => {
    if (!sel) return null;
    let els;
    try { els = [...document.querySelectorAll(sel)]; } catch (e) { return null; }
    return els.length ? els : null;
  };

  // окно распознанного текста: элемент справа от окна со сканом с самым большим текстом;
  // из равнозначных берётся самый вложенный (самый маленький по площади)
  const autoText = () => {
    const vr = findViewer();
    const minLeft = vr ? vr.right - 20 : window.innerWidth * 0.4;
    const skip = new Set(['SCRIPT', 'STYLE', 'CANVAS', 'IMG', 'SVG', 'NOSCRIPT', 'HEAD', 'HTML', 'BODY']);
    const cands = [];
    for (const e of document.body.querySelectorAll('*')) {
      if (skip.has(e.tagName.toUpperCase())) continue;
      const r = rectOf(e);
      if (r.width < 100 || r.height < 60 || r.left < minLeft) continue;
      if (!visible(e)) continue;
      const t = e.innerText || '';
      if (t.trim().length < 30) continue;
      cands.push({ e, len: t.length, area: r.width * r.height });
    }
    if (!cands.length) return null;
    const maxLen = Math.max(...cands.map((c) => c.len));
    const good = cands.filter((c) => c.len >= maxLen * 0.85).sort((a, b) => a.area - b.area);
    return { el: good[0].e, viewer: !!vr };
  };

  const getText = () => {
    const sel = bySelector(textSel);
    if (sel) return { text: sel.map((e) => e.innerText || '').filter((s) => s.trim()).join('\n'), how: 'вручную', found: true };
    const a = autoText();
    if (a) return { text: a.el.innerText || '', how: a.viewer ? 'авто' : 'авто (окно скана не найдено, взята правая часть)', found: true };
    return { text: '', how: '', found: false };
  };

  // дата документа: строка выше окна со сканом (над панелью инструментов), содержащая год
  const autoDate = () => {
    const vr = findViewer();
    if (!vr) return '';
    const cands = [];
    for (const e of document.body.querySelectorAll('*')) {
      const r = rectOf(e);
      if (r.width < 20 || r.height < 8 || r.bottom > vr.top + 4) continue;
      if (r.right < vr.left || r.left > vr.right) continue;
      if (vr.top - r.bottom > 300) continue;
      if (!visible(e)) continue;
      const t = clean(e.innerText);
      if (t.length < 4 || t.length > 200 || t.includes('\n')) continue;
      if (!/(?<!\d)1[5-9]\d\d(?!\d)/.test(t)) continue;
      cands.push({ e, t, bottom: r.bottom, score: (MONTHS.test(t) ? 2 : 0) + (/г\.|года|год/i.test(t) ? 1 : 0) });
    }
    // оставляем самые вложенные (без вложенных кандидатов)
    const leaf = cands.filter((c) => !cands.some((o) => o !== c && c.e.contains(o.e)));
    leaf.sort((a, b) => b.score - a.score || b.bottom - a.bottom);
    return leaf.length ? leaf[0].t : '';
  };

  const getDate = () => {
    const sel = bySelector(dateSel);
    if (sel) return { date: clean(sel.map((e) => e.innerText || '').join(' ')), how: 'вручную' };
    const d = autoDate();
    return { date: d, how: d ? 'авто' : '' };
  };

  // ждём, пока текст распознанной части прогрузится и перестанет меняться
  const t0 = Date.now();
  let last = -1, stable = 0, res = getText();
  while (Date.now() - t0 < maxWait) {
    res = getText();
    const len = res.text.trim().length;
    if (res.found && len > 0 && len === last) {
      if (++stable >= 3) break;
    } else {
      stable = 0;
    }
    if (res.found && len === 0 && Date.now() - t0 > 6000) break; // пустая страница (чистый лист)
    last = len;
    await new Promise((r) => setTimeout(r, 600));
  }
  const d = getDate();
  return {
    found: res.found,
    text: res.found ? res.text : '',
    textHow: res.how,
    date: d.date,
    dateHow: d.how,
    html: opts.wantHtml ? document.documentElement.outerHTML : null,
  };
}
"""


# ============================================================================
# Разбор
# ============================================================================
def _int(s):
    try:
        return int(s)
    except (TypeError, ValueError):
        return None


def parse_links(html_path, mode_id=DEFAULT_MODE):
    """Возвращает список словарей: url, year, file_id, p_from, p_to, note."""
    span = get_mode(mode_id)["default_span"]
    with open(html_path, "r", encoding="utf-8", errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    items, seen = [], set()
    for a in soup.find_all("a", href=True):
        url = a["href"].strip()
        u = urlparse(url)
        if u.hostname not in ARCHIVE_HOSTS or not u.path.startswith("/archive"):
            continue
        if url in seen:
            continue
        seen.add(url)

        # Год: текст всего абзаца — в списке бывают "разорванные" ссылки вида <a>186</a>4.
        holder = a.find_parent("p") or a
        m = re.search(r"(?<!\d)(1[5-9]\d\d)(?!\d)", holder.get_text("", strip=True))
        year = int(m.group(1)) if m else None

        q = parse_qs(u.query)
        note = ""
        file_id, p_from, p_to = None, None, None

        if u.path.startswith("/archive/search"):
            file_id = (q.get("fileId") or [None])[0]
            p_from = _int((q.get("sheet_page_from") or [""])[0])
            p_to = _int((q.get("sheet_page_to") or [""])[0])
        elif u.path.startswith("/archive/catalog/"):
            parts = [x for x in u.path.split("/") if x]
            if len(parts) >= 4:
                file_id = parts[2]
                p_from = p_to = _int(parts[3])
                note = "ссылка на одну страницу каталога, диапазона в ссылке нет"

        if p_from is not None and p_to is None:
            p_to = p_from + span
            note = f"нет sheet_page_to — взято from+{span}"
        if p_from is not None and p_to is not None and p_from > p_to:
            p_from, p_to = p_to, p_from

        items.append(dict(url=url, year=year, file_id=file_id, p_from=p_from, p_to=p_to, note=note))
    return items


def find_settlements(text, mode_id=DEFAULT_MODE):
    settlements = compile_mode(mode_id)["settlements"]
    t = norm(text)
    return {name for name, rx in settlements.items() if rx.search(t)}


def match_header(line, nxt="", mode_id=DEFAULT_MODE):
    """('own', имя) / ('other', имя) / None."""
    c_ = compile_mode(mode_id)
    cand = [line]
    # заголовок мог разорваться на две строки: "Деревня" / "Острогина".
    # Склеиваем ТОЛЬКО если текущая строка — одно слово-префикс.
    if nxt and c_["prefix_only"].match(line):
        cand.append(line + " " + nxt)
    for c in cand:
        if len(c) > HEADER_MAX_LEN:
            continue
        n = norm(c)
        for name, rx in c_["settlements"].items():
            if rx.search(n):
                return "own", name
        m = c_["other_header"].search(c)
        if m:
            return "other", m.group(1)
    return None


def extract_people(sheet_text, state, strictly_inside, mode_id=DEFAULT_MODE):
    """state: {'current','other','started','household'} — переносится между страницами."""
    mode = get_mode(mode_id)
    c_ = compile_mode(mode_id)
    lines = [s.strip() for s in re.split(r"\r?\n", sheet_text) if s.strip()]
    people = []
    skip_next = False
    for i, ln in enumerate(lines):
        if skip_next:                       # вторая половина разорванного заголовка
            skip_next = False
            continue
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        h = match_header(ln, nxt, mode_id)
        if h:
            kind, name = h
            state["current"] = name if kind == "own" else None
            state["other"] = name if kind == "other" else None
            state["household"] = None
            skip_next = bool(c_["prefix_only"].match(ln))
            continue
        if mode["household"]:
            hm = c_["hh_a"].match(ln) or c_["hh_b"].match(ln)
            if hm:
                state["household"] = int(hm.group(1))
                continue
        if len(ln) < mode["min_line_len"] or not re.search(r"[А-ЯЁІѢ][а-яёіѣъь]+", ln):
            continue
        if TITLE_NOISE.search(ln):          # служебные строки: название документа, уезд, приход...
            continue
        cur = state.get("current")
        if cur is None:
            # заголовка своего НП ещё не было: берём только внутри диапазона
            # и только если не встретился чужой НП
            if state.get("other") or not strictly_inside or not state.get("started"):
                continue
            cur = "не определён (начало диапазона)"
        ages = [int(x) for x in re.findall(r"(?<!\w)(\d{1,3})(?!\w)", ln)]
        ages = [x for x in ages if 0 < x < 111]
        people.append({
            "settlement": cur,
            "household": state["household"] if mode["household"] and state.get("household") else "",
            "line_no": i + 1,
            "ages": ", ".join(str(x) for x in ages),
            "age1": ages[0] if mode["two_ages"] and len(ages) > 0 else "",
            "age2": ages[1] if mode["two_ages"] and len(ages) > 1 else "",
            "text": ln,
        })
    return people


# ============================================================================
# Ввод/вывод
# ============================================================================
def ask_paths():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    html_path = filedialog.askopenfilename(
        title="Выберите HTML-файл со списком ссылок (list_links.html)",
        filetypes=[("HTML", "*.html *.htm"), ("Все файлы", "*.*")])
    if not html_path:
        sys.exit("Файл со ссылками не выбран.")
    xlsx_path = filedialog.asksaveasfilename(
        title="Куда сохранить результат (Excel)", defaultextension=".xlsx",
        initialfile="result.xlsx", filetypes=[("Excel", "*.xlsx")])
    if not xlsx_path:
        sys.exit("Путь сохранения не выбран.")
    root.destroy()
    return html_path, xlsx_path


def ask_mode():
    """Выбор типа обработки (в расширении — переключатель в разделе 1)."""
    ids = list(MODES)
    print("Тип обработки:")
    for i, mid in enumerate(ids, 1):
        print(f"  {i}. {MODES[mid]['label']} — {MODES[mid]['description']}")
    while True:
        s = input(f"Выберите номер [1-{len(ids)}] (Enter = {MODES[DEFAULT_MODE]['label']}): ").strip()
        if not s:
            return DEFAULT_MODE
        if s.isdigit() and 1 <= int(s) <= len(ids):
            return ids[int(s) - 1]


NEEDS_USER = re.compile(r"captcha|passport\.yandex|/auth", re.I)


def ensure_login(page):
    print("Проверяю вход в Яндекс...")
    try:
        page.goto(WORK_HOST + "/archive", wait_until="domcontentloaded", timeout=60_000)
    except Exception as e:  # noqa: BLE001
        print("Страница ya.ru/archive открылась с ошибкой:", e)
    on_archive = re.match(r"^https://(ya\.ru|yandex\.ru)/archive", page.url or "")
    if not on_archive or NEEDS_USER.search(page.url or ""):
        input("\n>>> Войдите в Яндекс в окне браузера (должен быть открыт ya.ru/archive) "
              "и нажмите Enter здесь...")
    # прогреваем сессию на втором домене (yandex.ru) — на случай ссылок с ним
    try:
        page.goto("https://yandex.ru/archive", wait_until="domcontentloaded", timeout=30_000)
        page.goto(WORK_HOST + "/archive", wait_until="domcontentloaded", timeout=30_000)
    except Exception:  # noqa: BLE001
        pass
    print("Вход проверен.")


_debug_saved = False


def fetch_sheet(page, file_id, n):
    """Открывает страницу n документа, возвращает dict: url, text, date, found, err."""
    global _debug_saved
    url = PAGE_URL_TEMPLATE.format(file_id=file_id, page=n)
    last_err = ""
    for attempt in range(MAX_RETRIES + 1):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
            guard = 0
            while NEEDS_USER.search(page.url or "") and guard < 3:
                guard += 1
                if "captcha" in page.url:
                    input("\nЯндекс просит капчу. Решите её в браузере и нажмите Enter...")
                else:
                    input("\nСессия Яндекса сбросилась. Войдите заново в браузере и нажмите Enter...")
                page.goto(url, wait_until="domcontentloaded", timeout=PAGE_TIMEOUT_MS)
            if (page.url or "").startswith("chrome-error"):
                raise RuntimeError("страница не загрузилась (ошибка сети)")

            want_html = DEBUG_HTML and not _debug_saved
            res = page.evaluate(EXTRACT_JS, {"textSel": TEXT_SELECTOR, "dateSel": DATE_SELECTOR,
                                             "maxWait": RENDER_WAIT_MS, "wantHtml": want_html}) or {}
            if want_html and res.get("html"):
                os.makedirs(DEBUG_DIR, exist_ok=True)
                with open(os.path.join(DEBUG_DIR, "first_page.html"), "w", encoding="utf-8") as f:
                    f.write(res["html"])
                _debug_saved = True
                print(f"Сохранён отладочный файл {DEBUG_DIR}/first_page.html")
            if res.get("found"):
                print(f"  стр. {n}: окно текста — {res.get('textHow')}, знаков: {len(res.get('text') or '')}"
                      + (f"; дата: «{res['date']}» ({res.get('dateHow')})" if res.get("date") else "; дата не найдена"))
            return {"url": url, "text": res.get("text") or "", "date": res.get("date") or "",
                    "found": bool(res.get("found")),
                    "err": "" if res.get("found") else "окно распознанного текста не найдено"}
        except Exception as e:  # noqa: BLE001
            last_err = str(e) or type(e).__name__
            print(f"  страница {n}: ошибка «{last_err}», попытка {attempt + 1}/{MAX_RETRIES + 1}")
            time.sleep(2)
    return {"url": url, "text": "", "date": "", "found": False, "err": last_err or "ошибка"}


def _c(v):
    v = "" if v is None else v
    return v[:32000] if isinstance(v, str) else v


def _widths(sheet):
    for col in sheet.columns:
        w = 8
        for x in col[:200]:
            if x.value not in (None, ""):
                w = max(w, len(str(x.value)))
        sheet.column_dimensions[col[0].column_letter].width = min(max(w + 2, 10), 70)


def save_workbook(path, mode_id, people, sheets, links):
    revision = get_mode(mode_id)["id"] == "revision"
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Люди"
    ws.append(["Год", "Тип документа", "Дата документа (со страницы)", "Населённый пункт", "Двор №",
               "Страница", "Строка на странице", "Числа-возрасты (черновик)",
               "Возраст: 1-е число" if revision else "", "Возраст: 2-е число" if revision else "",
               "Распознанная строка", "Ссылка на страницу"])
    for r in people:
        ws.append([_c(x) for x in (r["year"], r["doc_type"], r["doc_date"], r["settlement"], r["household"],
                                   r["page"], r["line_no"], r["ages"], r["age1"], r["age2"],
                                   r["text"], r["url"])])

    ws2 = wb.create_sheet("Страницы")
    ws2.append(["Год", "Тип документа", "Страница", "Дата документа (со страницы)",
                "Найденные НП округи", "Ссылка", "Распознанный текст страницы"])
    for r in sheets:
        ws2.append([_c(x) for x in (r["year"], r["doc_type"], r["page"], r["doc_date"],
                                    r["found"], r["url"], r["text"])])

    ws3 = wb.create_sheet("Ссылки")
    ws3.append(["Год", "Тип документа", "Исходная ссылка", "Диапазон из ссылки", "Просмотрено страниц",
                "Найдено НП (шт.)", "Какие НП", "Решение", "Людей собрано",
                "Дата документа (1-я страница)", "Примечание"])
    for r in links:
        ws3.append([_c(x) for x in r])

    for sh in (ws, ws2, ws3):
        sh.freeze_panes = "A2"
        for cell in sh[1]:
            cell.font = bold
        _widths(sh)
    wb.save(path)


def main():
    html_path, xlsx_path = ask_paths()
    mode_id = ask_mode()
    m = get_mode(mode_id)
    items = parse_links(html_path, mode_id)
    if not items:
        sys.exit("В файле не найдено ссылок Яндекс Архива.")
    print(f"Ссылок Яндекс Архива: {len(items)}. Тип обработки: {m['label']}.")

    people, sheet_rows, link_rows = [], [], []
    page_cache = {}          # (file_id, page) -> результат fetch_sheet
    seen_people = set()
    stop = False
    misses = 0

    with sync_playwright() as pw:
        # 1) пробуем подключиться к уже запущенному Chrome (порт 9222), 2) иначе запускаем установленный Chrome
        try:
            browser = pw.chromium.connect_over_cdp("http://localhost:9222")
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            print("Подключился к запущенному Chrome (порт 9222).")
        except Exception:  # noqa: BLE001
            print("Chrome с портом 9222 не найден — запускаю установленный Google Chrome...")
            browser = pw.chromium.launch(headless=False, channel="chrome")
            context = browser.new_context(locale="ru-RU")
        page = context.new_page()

        try:
            ensure_login(page)
            print(f"Окно текста: {TEXT_SELECTOR or 'авто'}; дата: {DATE_SELECTOR or 'авто'}.")

            for idx, it in enumerate(items, 1):
                if stop:
                    break
                year = it["year"]
                tag = f"[{idx}/{len(items)}] {year if year is not None else '—'}"
                rng = f"{it['p_from']}–{it['p_to']}" if it["p_from"] is not None else "—"

                if not it["file_id"] or it["p_from"] is None:
                    link_rows.append([year, m["label"], it["url"], rng, 0, 0, "",
                                      "пропущена: нет fileId или диапазона", 0, "", it["note"]])
                    print(f"{tag}: пропущена (нет fileId/диапазона)")
                    continue

                lo, hi = it["p_from"], it["p_to"]
                if hi - lo + 1 > BIG_RANGE:
                    print(f"{tag}: внимание, диапазон очень большой ({hi - lo + 1} страниц)")

                # 1. загрузка страниц диапазона из ссылки (диапазон — ровно из ссылки)
                texts = {}
                try:
                    for n in range(lo, hi + 1):
                        if stop:
                            break
                        key = (it["file_id"], n)
                        if key not in page_cache:
                            print(f"{tag}: загружаю страницу {n} (диапазон {lo}–{hi})...")
                            r = fetch_sheet(page, it["file_id"], n)
                            page_cache[key] = r
                            if not r["found"]:
                                misses += 1
                                if misses >= MAX_MISSES:
                                    stop = True
                                    print(f"Окно распознанного текста не найдено на {MAX_MISSES} страницах подряд — "
                                          "останавливаюсь, чтобы не собирать лишнее. Задайте TEXT_SELECTOR "
                                          "в начале скрипта и запустите снова.")
                            else:
                                misses = 0
                            if n < hi:
                                time.sleep(random.uniform(*DELAY_RANGE))
                        texts[n] = page_cache[key]
                except KeyboardInterrupt:
                    print("\nОстановка запрошена — сохраняю собранное...")
                    stop = True

                # 2. сбор людей
                found_all, per_page = set(), {}
                for n, r in texts.items():
                    f = find_settlements(r["text"], mode_id)
                    per_page[n] = f
                    found_all |= f

                state = {"current": None, "other": None, "started": True, "household": None}
                collected = 0
                first_date = ""
                for n in sorted(texts):
                    r = texts[n]
                    if not first_date and r["date"]:
                        first_date = r["date"]
                    sheet_rows.append({"year": year, "doc_type": m["label"], "page": n, "doc_date": r["date"],
                                       "found": "; ".join(sorted(per_page[n])), "url": r["url"],
                                       "text": f"ОШИБКА: {r['err']}" if r["err"] else r["text"]})
                    for p in extract_people(r["text"], state, True, mode_id):
                        key = (it["file_id"], n, p["line_no"], p["text"])
                        if key in seen_people:
                            continue
                        seen_people.add(key)
                        p.update(year=year, doc_type=m["label"], doc_date=r["date"], page=n, url=r["url"])
                        people.append(p)
                        collected += 1

                total = hi - lo + 1
                verdict = f"ПРЕРВАНА (страниц {len(texts)} из {total})" if len(texts) < total else "ОБРАБОТАНА"
                link_rows.append([year, m["label"], it["url"], rng, len(texts), len(found_all),
                                  "; ".join(sorted(found_all)), verdict, collected, first_date, it["note"]])
                print(f"{tag}: диапазон {rng} — НП {len(found_all)}, людей: {collected}")
                save_workbook(xlsx_path, mode_id, people, sheet_rows, link_rows)
        except KeyboardInterrupt:
            print("\nОстановлено пользователем.")
        finally:
            try:
                save_workbook(xlsx_path, mode_id, people, sheet_rows, link_rows)
            finally:
                try:
                    browser.close()
                except Exception:  # noqa: BLE001
                    pass

    print(f"\nГотово. Людей: {len(people)}. Файл: {xlsx_path}")
    try:
        r = tk.Tk()
        r.withdraw()
        messagebox.showinfo("Готово", f"Сохранено: {xlsx_path}\nЛюдей: {len(people)}")
    except Exception:  # noqa: BLE001
        pass


if __name__ == "__main__":
    main()
